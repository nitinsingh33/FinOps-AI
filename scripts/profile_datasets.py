"""
Dataset profiling script — reads all 4 Excel workbooks and prints
sheet names, header rows, column names, dtypes, record counts, sample values.
"""
import sys
import json
import openpyxl
import pandas as pd
from pathlib import Path

DATA_DIR = Path(__file__).parent.parent / "data"

FILES = [
    "GatewayA_payments_Sep26_dummy.xlsx",
    "GatewayA_transactions_Sep26_dummy.xlsx",
    "GatewayA_refunds_Sep26_dummy.xlsx",
    "GatewayB_Sep26_dummy.xlsx",
]

def find_header_row(ws, max_scan=10):
    """Scan first max_scan rows to find the real header row (most non-blank cells)."""
    best_row = 0
    best_count = 0
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=max_scan, values_only=True)):
        count = sum(1 for c in row if c is not None and str(c).strip())
        if count > best_count:
            best_count = count
            best_row = i + 1  # 1-indexed
    return best_row


def profile_file(filepath: Path):
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    report = {"file": filepath.name, "sheets": []}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        header_row_idx = find_header_row(ws)
        # Re-read with pandas using detected header row
        df = pd.read_excel(
            filepath,
            sheet_name=sheet_name,
            header=header_row_idx - 1,  # 0-indexed for pandas
            engine="openpyxl",
        )
        df.columns = [str(c).strip() for c in df.columns]
        # Drop rows that are completely empty
        df = df.dropna(how="all")
        sheet_report = {
            "sheet": sheet_name,
            "detected_header_row": header_row_idx,
            "total_rows_after_header": len(df),
            "columns": [],
        }
        for col in df.columns:
            series = df[col]
            non_null = series.dropna()
            sheet_report["columns"].append({
                "name": col,
                "dtype": str(series.dtype),
                "non_null_count": int(non_null.count()),
                "null_count": int(series.isna().sum()),
                "null_pct": round(series.isna().mean() * 100, 1),
                "unique_count": int(non_null.nunique()),
                "sample_values": [str(v) for v in non_null.head(3).tolist()],
            })
        report["sheets"].append(sheet_report)
    wb.close()
    return report


if __name__ == "__main__":
    all_reports = []
    for fname in FILES:
        fpath = DATA_DIR / fname
        if not fpath.exists():
            print(f"MISSING: {fpath}", file=sys.stderr)
            continue
        print(f"\n{'='*60}")
        print(f"Profiling: {fname}")
        print(f"{'='*60}")
        report = profile_file(fpath)
        all_reports.append(report)
        for sheet in report["sheets"]:
            print(f"\n  Sheet: '{sheet['sheet']}' | Header row: {sheet['detected_header_row']} | Rows: {sheet['total_rows_after_header']}")
            for col in sheet["columns"]:
                print(f"    {col['name']!r:40s} dtype={col['dtype']:12s} nulls={col['null_pct']:5.1f}%  unique={col['unique_count']:6d}  sample={col['sample_values'][:2]}")
    
    out_path = DATA_DIR / "profile_report.json"
    with open(out_path, "w") as f:
        json.dump(all_reports, f, indent=2, default=str)
    print(f"\n\nFull JSON report written to: {out_path}")
